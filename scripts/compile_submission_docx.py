from __future__ import annotations

import os
import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def set_cell_shading(cell, color_hex: str) -> None:
    """Apply background color shading to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150) -> None:
    """Set padding/margins for a table cell (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_formatted_text(paragraph, text: str) -> None:
    """Parse basic inline markdown (*italic*, **bold**, `code`) and add it to paragraph."""
    # Find patterns of bold, italic, and inline code
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def parse_markdown_table(lines: list[str]) -> list[list[str]]:
    """Parse markdown table lines into a list of row values."""
    rows = []
    for line in lines:
        if not line.strip() or re.match(r'^\|?\s*:?-+:?\s*(\|?\s*:?-+:?\s*)*$', line.strip()):
            # Skip separator line
            continue
        # Split by | and strip whitespace
        parts = [p.strip() for p in line.split('|')]
        # Handle leading/trailing |
        if parts and parts[0] == '':
            parts = parts[1:]
        if parts and parts[-1] == '':
            parts = parts[:-1]
        rows.append(parts)
    return rows


def compile_md_to_docx(md_path: Path, docx_path: Path, title: str) -> None:
    print(f"Compiling {md_path.name} to {docx_path.name}...")
    doc = docx.Document()

    # Set standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Set normal font to Times New Roman, 11 pt
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    # Paragraph spacing
    p_format = style.paragraph_format
    p_format.line_spacing = 1.15
    p_format.space_after = Pt(6)

    # Read markdown file
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Parse and write to docx
    table_buffer: list[str] = []
    in_table = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle table state
        if stripped.startswith('|'):
            in_table = True
            table_buffer.append(line)
            i += 1
            continue
        elif in_table:
            # End of table block
            in_table = False
            table_data = parse_markdown_table(table_buffer)
            if table_data:
                num_cols = len(table_data[0])
                table = doc.add_table(rows=len(table_data), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Table Grid'
                
                for r_idx, row_values in enumerate(table_data):
                    row = table.rows[r_idx]
                    # Make header row bold and shaded
                    is_header = (r_idx == 0)
                    for c_idx, val in enumerate(row_values):
                        if c_idx < len(row.cells):
                            cell = row.cells[c_idx]
                            cell.text = "" # Clear default text
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.0
                            
                            # Alignment: right align columns if they seem numeric/monetary
                            if r_idx > 0 and (val.startswith('$') or val.endswith('%') or re.match(r'^\d+(\.\d+)?$', val)):
                                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            
                            run = p.add_run(val)
                            if is_header:
                                run.bold = True
                                set_cell_shading(cell, 'F2F2F2')
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            table_buffer = []

        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.size = Pt(16)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:])
            run.bold = True
            run.font.size = Pt(13)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.size = Pt(11.5)
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.size = Pt(11)
        # Bullet list items
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, stripped[2:])
        # Numbered list items
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^(\d+\.\s)(.*)', stripped)
            p = doc.add_paragraph()
            # Use manual left indent for list look
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, stripped)
        # Standard paragraph
        else:
            p = doc.add_paragraph()
            add_formatted_text(p, stripped)

        i += 1

    # Save document
    doc.save(docx_path)
    print(f"Successfully saved {docx_path.name}")


def main() -> None:
    # Set paths
    project_root = Path(__file__).resolve().parent.parent
    submission_dir = project_root / 'docs' / 'submission'
    
    # 1. Compile Volume 2 Technical Proposal
    v2_md = submission_dir / 'volume_2_10_page_submission_draft.md'
    v2_docx = submission_dir / 'volume_2_10_page_submission_draft.docx'
    if v2_md.exists():
        compile_md_to_docx(v2_md, v2_docx, "Volume 2 Technical Proposal")
    else:
        print(f"File not found: {v2_md}")
        
    # 2. Compile Volume 3 Cost Proposal
    v3_md = submission_dir / 'volume_3_cost_draft.md'
    v3_docx = submission_dir / 'volume_3_cost_draft.docx'
    if v3_md.exists():
        compile_md_to_docx(v3_md, v3_docx, "Volume 3 Cost Proposal")
    else:
        print(f"File not found: {v3_md}")

    # 3. Compile Volume 5 Supporting Documents
    v5_md = submission_dir / 'volume_5_supporting_docs_draft.md'
    v5_docx = submission_dir / 'volume_5_supporting_docs_draft.docx'
    if v5_md.exists():
        compile_md_to_docx(v5_md, v5_docx, "Volume 5 Supporting Documents")
    else:
        print(f"File not found: {v5_md}")


if __name__ == '__main__':
    main()
